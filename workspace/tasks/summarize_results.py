#!/usr/bin/env python3
"""Summarize finished (task, skill) pairs into the chapter-4 Excel workbook.

Usage: summarize_results.py <out_xlsx> <run_dir> [<run_dir> ...]

Each run_dir is one finished run-skill-batch.sh output folder
(results/<task>-<skill>/run-<N>/) containing runs-tokens.csv, report-KEY.md
and report-scored.md. Per pair and condition (off / skill) it reports the
median and min-max of output tokens, input-side tokens, cost, cold-equivalent
cost and checklist score, then applies the decision rule from chapter 3.10:
ranges that do not overlap = a real difference, overlapping = inconclusive.
"""
import csv, os, re, statistics, sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def read_pair(run_dir):
    rows = list(csv.DictReader(open(os.path.join(run_dir, "runs-tokens.csv"), encoding="utf-8")))
    key = open(os.path.join(run_dir, "report-KEY.md"), encoding="utf-8").read()
    label_to_code = dict(re.findall(r"^\| ([A-Z]) \| (\S+) \|", key, re.M))
    scored = open(os.path.join(run_dir, "report-scored.md"), encoding="utf-8").read()
    scores = {label_to_code[l]: float(p) for l, p in
              re.findall(r"^## Run ([A-Z]) .*?\(([\d.]+)% passed\)", scored, re.M)}
    verdicts = {label_to_code[l]: v for l, v in re.findall(r"^## Run ([A-Z]) .*?(PASS|FAIL)", scored, re.M)}
    out = []
    for r in rows:
        code = r["run_code"]
        out.append({
            "task": r["task_id"], "skill": r["skill_name"], "cond": r["condition"], "code": code,
            "output": int(r["output"]),
            "input_side": int(r["input"]) + int(r["cache_read"]) + int(r["cache_w5"]) + int(r["cache_w1h"]),
            "cost": float(r["cost_usd"]), "cold": float(r["cost_cold_usd"]),
            "score": scores.get(code), "verdict": verdicts.get(code),
        })
    return out


def stats(vals):
    vals = [v for v in vals if v is not None]
    if not vals:
        return None
    return statistics.median(vals), min(vals), max(vals)


def decide(off, skill, lower_is_better):
    if off is None or skill is None:
        return "ไม่มีข้อมูล"
    if off[2] < skill[1] or skill[2] < off[1]:
        skill_lower = skill[2] < off[1]
        if lower_is_better:
            return "skill ดีกว่า (ต่ำกว่า)" if skill_lower else "skill แย่กว่า (สูงกว่า)"
        return "skill แย่กว่า (ต่ำกว่า)" if skill_lower else "skill ดีกว่า (สูงกว่า)"
    return "สรุปไม่ได้ (ช่วงทับกัน)"


def fmt(st, nd=0):
    if st is None:
        return "-"
    f = (lambda x: f"{x:,.{nd}f}")
    return f"{f(st[0])} ({f(st[1])} - {f(st[2])})"


METRICS = [("output", "Output Token", 0, True), ("input_side", "Input Token รวมแคช", 0, True),
           ("cost", "ต้นทุน (USD)", 4, True), ("cold", "Cold-equivalent Cost (USD)", 4, True),
           ("score", "คะแนนความถูกต้อง (%)", 1, False)]


LABELS = {"caveman": "Caveman",
          "verification-before-completion": "verification (ให้โมเดลเลือกเอง)",
          "verification-forced": "verification (สั่งเรียกใช้)"}


def main():
    out_xlsx, run_dirs = sys.argv[1], sys.argv[2:]
    pairs = [read_pair(d) for d in run_dirs]
    wb = Workbook(); wb.remove(wb.active)
    thin = Side(style="thin", color="000000"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    f = Font(name="Angsana New", size=14); fb = Font(name="Angsana New", size=14, bold=True)
    fill = PatternFill("solid", fgColor="D9E2F3")

    def sheet(name, title, header, rows):
        ws = wb.create_sheet(name); ws["A1"] = title; ws["A1"].font = fb
        for r, row in enumerate([header] + rows):
            for c, v in enumerate(row, 1):
                cell = ws.cell(r + 3, c, v); cell.font = fb if r == 0 else f; cell.border = bd
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if r == 0: cell.fill = fill
        for i in range(1, len(header) + 1):
            ws.column_dimensions[chr(64 + i)].width = 24
        return ws

    t3, t4, t5, chart_rows = [], [], [], []
    for runs in pairs:
        task, skill = runs[0]["task"], LABELS.get(runs[0]["skill"], runs[0]["skill"])
        by = {c: [r for r in runs if r["cond"] == c] for c in ("off", "skill")}
        st = {(c, m): stats([r[m] for r in by[c]]) for c in by for m, *_ in METRICS}
        t3.append([task, skill] + [fmt(st[(c, m)], nd) for m, _, nd, _ in METRICS[:4] for c in ("off", "skill")])
        passes = {c: sum(1 for r in by[c] if r["verdict"] == "PASS") for c in by}
        t4.append([task, skill, fmt(st[("off", "score")], 1), fmt(st[("skill", "score")], 1),
                   f"{passes['off']}/{len(by['off'])}", f"{passes['skill']}/{len(by['skill'])}"])
        t5.append([task, skill] + [decide(st[("off", m)], st[("skill", m)], low) for m, _, _, low in METRICS])
        chart_rows.append([f"{task}\n{skill}", st[("off", "cost")][0], st[("skill", "cost")][0],
                           st[("off", "output")][0], st[("skill", "output")][0]])

    hdr3 = ["งาน", "เครื่องมือ"] + [f"{lbl} {c}" for _, lbl, _, _ in METRICS[:4] for c in ("off", "skill")]
    sheet("4.3 Token ต้นทุน", "ตารางที่ 4.3 ปริมาณ Token และต้นทุน ค่ามัธยฐาน (ต่ำสุด - สูงสุด) จาก 3 รอบ", hdr3, t3)
    sheet("4.4 ความถูกต้อง", "ตารางที่ 4.4 คะแนนความถูกต้อง ค่ามัธยฐาน (ต่ำสุด - สูงสุด) และจำนวนรอบที่ผ่านเกณฑ์ร้อยละ 80",
          ["งาน", "เครื่องมือ", "คะแนน off", "คะแนน skill", "รอบที่ผ่าน off", "รอบที่ผ่าน skill"], t4)
    sheet("4.5 สรุปผล", "ตารางที่ 4.5 ผลการตัดสินตามกฎช่วงค่าไม่ทับซ้อนกัน",
          ["งาน", "เครื่องมือ"] + [lbl for _, lbl, _, _ in METRICS], t5)

    ws = sheet("กราฟ", "ข้อมูลกราฟ (ค่ามัธยฐาน)",
               ["คู่", "ต้นทุน off", "ต้นทุน skill", "Output off", "Output skill"], chart_rows)
    n = len(chart_rows)
    for col0, title, anchor in ((2, "รูปที่ 4.1 ต้นทุนมัธยฐานต่อรอบ (USD)", "H3"),
                                (4, "รูปที่ 4.2 Output Token มัธยฐานต่อรอบ", "H22")):
        ch = BarChart(); ch.type = "col"; ch.width = 22; ch.height = 9
        # openpyxl 3.1 hides both axes by default; the figure caption lives in the report, not the chart
        ch.x_axis.delete = False; ch.y_axis.delete = False
        ch.y_axis.title = "USD" if col0 == 2 else "Token"
        ch.y_axis.majorGridlines = None
        ch.add_data(Reference(ws, min_col=col0, max_col=col0 + 1, min_row=3, max_row=3 + n), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + n))
        ws.add_chart(ch, anchor)

    raw = [[r["task"], LABELS.get(r["skill"], r["skill"]), r["cond"], r["code"], r["output"], r["input_side"],
            round(r["cost"], 6), round(r["cold"], 6), r["score"], r["verdict"]] for runs in pairs for r in runs]
    sheet("ข้อมูลดิบ", "ข้อมูลดิบทุกรอบที่นับผล",
          ["งาน", "เครื่องมือ", "เงื่อนไข", "รหัสรอบ", "Output", "Input รวมแคช", "Cost", "Cold Cost", "คะแนน %", "ผล"], raw)
    wb.save(out_xlsx)
    print(f"wrote {out_xlsx}: {len(pairs)} pairs")


if __name__ == "__main__":
    main()
